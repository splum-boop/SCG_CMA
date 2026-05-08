"""Step 15b: Counties Reference — states 11-20 (Hawaii through Maryland)."""
import openpyxl
from openpyxl.styles import Font, Border, Side

def thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

counties = {
    'Hawaii': 'Hawaii,Honolulu,Kalawao,Kauai,Maui',
    'Idaho': 'Ada,Adams,Bannock,Bear Lake,Benewah,Bingham,Blaine,Boise,Bonner,Bonneville,Boundary,Butte,Camas,Canyon,Caribou,Cassia,Clark,Clearwater,Custer,Elmore,Franklin,Fremont,Gem,Gooding,Idaho,Jefferson,Jerome,Kootenai,Latah,Lemhi,Lewis,Lincoln,Madison,Minidoka,Nez Perce,Oneida,Owyhee,Payette,Power,Shoshone,Teton,Twin Falls,Valley,Washington',
    'Illinois': 'Adams,Alexander,Bond,Boone,Brown,Bureau,Calhoun,Carroll,Cass,Champaign,Christian,Clark,Clay,Clinton,Coles,Cook,Crawford,Cumberland,DeKalb,De Witt,Douglas,DuPage,Edgar,Edwards,Effingham,Fayette,Ford,Franklin,Fulton,Gallatin,Greene,Grundy,Hamilton,Hancock,Hardin,Henderson,Henry,Iroquois,Jackson,Jasper,Jefferson,Jersey,Jo Daviess,Johnson,Kane,Kankakee,Kendall,Knox,Lake,LaSalle,Lawrence,Lee,Livingston,Logan,Macon,Macoupin,Madison,Marion,Marshall,Mason,Massac,McDonough,McHenry,McLean,Menard,Mercer,Monroe,Montgomery,Morgan,Moultrie,Ogle,Peoria,Perry,Piatt,Pike,Pope,Pulaski,Putnam,Randolph,Richland,Rock Island,Saline,Sangamon,Schuyler,Scott,Shelby,St. Clair,Stark,Stephenson,Tazewell,Union,Vermilion,Wabash,Warren,Washington,Wayne,White,Whiteside,Will,Williamson,Winnebago,Woodford',
    'Indiana': 'Adams,Allen,Bartholomew,Benton,Blackford,Boone,Brown,Carroll,Cass,Clark,Clay,Clinton,Crawford,Daviess,Dearborn,Decatur,DeKalb,Delaware,Dubois,Elkhart,Fayette,Floyd,Fountain,Franklin,Fulton,Gibson,Grant,Greene,Hamilton,Hancock,Harrison,Hendricks,Henry,Howard,Huntington,Jackson,Jasper,Jay,Jefferson,Jennings,Johnson,Knox,Kosciusko,LaGrange,Lake,LaPorte,Lawrence,Madison,Marion,Marshall,Martin,Miami,Monroe,Montgomery,Morgan,Newton,Noble,Ohio,Orange,Owen,Parke,Perry,Pike,Porter,Posey,Pulaski,Putnam,Randolph,Ripley,Rush,Scott,Shelby,Spencer,St. Joseph,Starke,Steuben,Sullivan,Switzerland,Tippecanoe,Tipton,Union,Vanderburgh,Vermillion,Vigo,Wabash,Warren,Warrick,Washington,Wayne,Wells,White,Whitley',
    'Iowa': "Adair,Adams,Allamakee,Appanoose,Audubon,Benton,Black Hawk,Boone,Bremer,Buchanan,Buena Vista,Butler,Calhoun,Carroll,Cass,Cedar,Cerro Gordo,Cherokee,Chickasaw,Clarke,Clay,Clayton,Clinton,Crawford,Dallas,Davis,Decatur,Delaware,Des Moines,Dickinson,Dubuque,Emmet,Fayette,Floyd,Franklin,Fremont,Greene,Grundy,Guthrie,Hamilton,Hancock,Hardin,Harrison,Henry,Howard,Humboldt,Ida,Iowa,Jackson,Jasper,Jefferson,Johnson,Jones,Keokuk,Kossuth,Lee,Linn,Louisa,Lucas,Lyon,Madison,Mahaska,Marion,Marshall,Mills,Mitchell,Monona,Monroe,Montgomery,Muscatine,O'Brien,Osceola,Page,Palo Alto,Plymouth,Pocahontas,Polk,Pottawattamie,Poweshiek,Ringgold,Sac,Scott,Shelby,Sioux,Story,Tama,Taylor,Union,Van Buren,Wapello,Warren,Washington,Wayne,Webster,Winnebago,Winneshiek,Woodbury,Worth,Wright",
    'Kansas': 'Allen,Anderson,Atchison,Barber,Barton,Bourbon,Brown,Butler,Chase,Chautauqua,Cherokee,Cheyenne,Clark,Clay,Cloud,Coffey,Comanche,Cowley,Crawford,Decatur,Dickinson,Doniphan,Douglas,Edwards,Elk,Ellis,Ellsworth,Finney,Ford,Franklin,Geary,Gove,Graham,Grant,Gray,Greeley,Greenwood,Hamilton,Harper,Harvey,Haskell,Hodgeman,Jackson,Jefferson,Jewell,Johnson,Kearny,Kingman,Kiowa,Labette,Lane,Leavenworth,Lincoln,Linn,Logan,Lyon,Marion,Marshall,McPherson,Meade,Miami,Mitchell,Montgomery,Morris,Morton,Nemaha,Neosho,Ness,Norton,Osage,Osborne,Ottawa,Pawnee,Phillips,Pottawatomie,Pratt,Rawlins,Reno,Republic,Rice,Riley,Rooks,Rush,Russell,Saline,Scott,Sedgwick,Seward,Shawnee,Sheridan,Sherman,Smith,Stafford,Stanton,Stevens,Sumner,Thomas,Trego,Wabaunsee,Wallace,Washington,Wichita,Wilson,Woodson,Wyandotte',
    'Kentucky': 'Adair,Allen,Anderson,Ballard,Barren,Bath,Bell,Boone,Bourbon,Boyd,Boyle,Bracken,Breathitt,Breckinridge,Bullitt,Butler,Caldwell,Calloway,Campbell,Carlisle,Carroll,Carter,Casey,Christian,Clark,Clay,Clinton,Crittenden,Cumberland,Daviess,Edmonson,Elliott,Estill,Fayette,Fleming,Floyd,Franklin,Fulton,Gallatin,Garrard,Grant,Graves,Grayson,Green,Greenup,Hancock,Hardin,Harlan,Harrison,Hart,Henderson,Henry,Hickman,Hopkins,Jackson,Jefferson,Jessamine,Johnson,Kenton,Knott,Knox,Larue,Laurel,Lawrence,Lee,Leslie,Letcher,Lewis,Lincoln,Livingston,Logan,Lyon,Madison,Magoffin,Marion,Marshall,Martin,Mason,McCracken,McCreary,McLean,Meade,Menifee,Mercer,Metcalfe,Monroe,Montgomery,Morgan,Muhlenberg,Nelson,Nicholas,Ohio,Oldham,Owen,Owsley,Pendleton,Perry,Pike,Powell,Pulaski,Robertson,Rockcastle,Rowan,Russell,Scott,Shelby,Simpson,Spencer,Taylor,Todd,Trigg,Trimble,Union,Warren,Washington,Wayne,Webster,Whitley,Wolfe,Woodford',
    'Louisiana': 'Acadia,Allen,Ascension,Assumption,Avoyelles,Beauregard,Bienville,Bossier,Caddo,Calcasieu,Caldwell,Cameron,Catahoula,Claiborne,Concordia,De Soto,East Baton Rouge,East Carroll,East Feliciana,Evangeline,Franklin,Grant,Iberia,Iberville,Jackson,Jefferson,Jefferson Davis,La Salle,Lafayette,Lafourche,Lincoln,Livingston,Madison,Morehouse,Natchitoches,Orleans,Ouachita,Plaquemines,Pointe Coupee,Rapides,Red River,Richland,Sabine,St. Bernard,St. Charles,St. Helena,St. James,St. John the Baptist,St. Landry,St. Martin,St. Mary,St. Tammany,Tangipahoa,Tensas,Terrebonne,Union,Vermilion,Vernon,Washington,Webster,West Baton Rouge,West Carroll,West Feliciana,Winn',
    'Maine': 'Androscoggin,Aroostook,Cumberland,Franklin,Hancock,Kennebec,Knox,Lincoln,Oxford,Penobscot,Piscataquis,Sagadahoc,Somerset,Waldo,Washington,York',
    'Maryland': 'Allegany,Anne Arundel,Baltimore,Baltimore City,Calvert,Caroline,Carroll,Cecil,Charles,Dorchester,Frederick,Garrett,Harford,Howard,Kent,Montgomery,Prince Georges,Queen Annes,Somerset,St. Marys,Talbot,Washington,Wicomico,Worcester',
}

wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Counties Reference"]
ws.sheet_state = "hidden"

# Find next empty row after existing data
next_row = 3
while ws[f"A{next_row}"].value:
    next_row += 1

for state, county_list in counties.items():
    ws[f"A{next_row}"].value = state
    ws[f"B{next_row}"].value = county_list
    ws[f"A{next_row}"].font  = Font(size=9, name="Arial")
    ws[f"B{next_row}"].font  = Font(size=9, name="Arial")
    ws[f"A{next_row}"].border = thin()
    ws[f"B{next_row}"].border = thin()
    next_row += 1

wb.save("SCG_CMA_System_v7.xlsx")

last_data_row = next_row - 1
print(f"Step 15b complete — {len(counties)} states appended (rows 13-{last_data_row}).")
print(f"States written: {', '.join(counties.keys())}")
print(f"Total data rows so far: {last_data_row - 2}")  # minus title + header
