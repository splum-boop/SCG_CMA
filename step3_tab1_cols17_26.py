"""Step 3: Tab 1 — cols 17-26 (Q-Z), YEAR LAST UPDATED section."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

# ── Colors ────────────────────────────────────────────────────────────────────
TEAL        = "1F6B75"
INSTRUCT_BG = "EAF0F6"
MED_GREY    = "CCCCCC"
DARK_GREY   = "888888"
WHITE       = "FFFFFF"
GREEN_BG    = "E2EFDA"
GREEN_TXT   = "375623"
YELLOW_BG   = "FFF2CC"
OLD_TXT     = "7D4E00"

# ── Helpers ───────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def apply(cell, value=None, fill=None, font=None, alignment=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    cell.border = thin_border()

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Column definitions (col index → letter, label, instruction, width) ────────
# Col 17=Q, 18=R, 19=S, 20=T, 21=U, 22=V, 23=W, 24=X, 25=Y, 26=Z
cols = [
    ("Q", "Roof",         "Roof — Year Last Replaced",    "Enter year or leave blank if unknown",           14),
    ("R", "Kitchen",      "Kitchen — Year Last Updated",  "Enter year or leave blank if unknown",           14),
    ("S", "Bathrooms",    "Bathrooms — Year Last Updated","Enter year or leave blank if unknown",           14),
    ("T", "HVAC",         "HVAC — Year Last Replaced",    "Enter year or leave blank if unknown",           14),
    ("U", "Windows",      "Windows — Year Last Replaced", "Enter year or leave blank if unknown",           14),
    ("V", "Electrical",   "Electrical — Year Last Updated","Enter year or leave blank if unknown",          14),
    ("W", "Plumbing",     "Plumbing — Year Last Updated", "Enter year or leave blank if unknown",           14),
    ("X", "Flooring",     "Flooring — Year Last Updated", "Enter year or leave blank if unknown",           14),
    ("Y", "Additions",    "Additions — Year Last Added",  "Year of most recent addition or major remodel",  14),
    ("Z", "Other Updates","Other Updates",                "Describe and include year",                       24),
]

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Property Submissions"]

# ── ROW 2 — Section Banner: Q2:Z2 ────────────────────────────────────────────
ws.merge_cells("Q2:Z2")
apply(ws["Q2"],
      value="YEAR LAST UPDATED",
      fill=solid(TEAL),
      font=Font(bold=True, size=9, color=WHITE, name="Arial"),
      alignment=center())

# ── ROW 3 — Headers ───────────────────────────────────────────────────────────
for col, _, label, _, _ in cols:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(TEAL),
          font=Font(bold=True, size=9, color=WHITE, name="Arial"),
          alignment=center(wrap=True))

# ── ROW 4 — Instructions ──────────────────────────────────────────────────────
for col, _, _, instruction, _ in cols:
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=Font(italic=True, size=7, color=DARK_GREY, name="Arial"),
          alignment=left_wrap())

# ── Column Widths ─────────────────────────────────────────────────────────────
for col, _, _, _, width in cols:
    ws.column_dimensions[col].width = width

# ── Conditional Formatting — cols 17-25 (Q-Y), data rows 5-2000 ───────────────
# GREEN when >= 2015
green_fill = solid(GREEN_BG)
green_font = Font(color=GREEN_TXT, name="Arial")

# YELLOW when < 2000
yellow_fill = solid(YELLOW_BG)
old_font    = Font(color=OLD_TXT, name="Arial")

for col in ["Q", "R", "S", "T", "U", "V", "W", "X", "Y"]:
    rng = f"{col}5:{col}2000"
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="greaterThanOrEqual",
        formula=["2015"],
        fill=green_fill,
        font=green_font,
    ))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan",
        formula=["2000"],
        fill=yellow_fill,
        font=old_font,
    ))

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 3 complete — Tab 1 cols 17-26 (Q-Z) built.")
print("  Section banner Q2:Z2 (TEAL), headers row 3, instructions row 4.")
print("  Conditional formatting on Q-Y: green >= 2015, yellow < 2000.")
