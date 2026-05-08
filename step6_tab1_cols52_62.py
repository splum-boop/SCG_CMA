"""Step 6: Tab 1 — cols 52-62 (AZ-BJ), Pending/Contract Tracking + Claude Output."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY_LIGHT  = "2E5F8A"
ORANGE_BG   = "FCE9D6"
INSTRUCT_BG = "EAF0F6"
MED_GREY    = "CCCCCC"
DARK_GREY   = "888888"
WHITE       = "FFFFFF"
GREEN_BG    = "E2EFDA"

# ── Helpers ───────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def apply(cell, value=None, fill=None, font=None, alignment=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    cell.border = thin_border()

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def hdr_font(color=WHITE):
    return Font(bold=True, size=9, color=color, name="Arial")

def instr_font():
    return Font(italic=True, size=7, color=DARK_GREY, name="Arial")

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Property Submissions"]

# ── Column map ────────────────────────────────────────────────────────────────
# Col 52=AZ, 53=BA
# Col 54=BB, 55=BC, 56=BD, 57=BE, 58=BF, 59=BG, 60=BH, 61=BI, 62=BJ

# ── ROW 2 — Section Banners ───────────────────────────────────────────────────
ws.merge_cells("AZ2:BA2")
apply(ws["AZ2"],
      value="PENDING / CONTRACT TRACKING",
      fill=solid(ORANGE_BG),
      font=hdr_font(color="000000"),
      alignment=center())

ws.merge_cells("BB2:BJ2")
apply(ws["BB2"],
      value="CLAUDE OUTPUT — Auto-populated by system",
      fill=solid(NAVY_LIGHT),
      font=hdr_font(),
      alignment=center())

# ── PENDING / CONTRACT TRACKING (AZ-BA) ──────────────────────────────────────
pending_cols = [
    ("AZ", "Contract Price",
     "Enter when property goes under contract",
     14, "$#,##0"),
    ("BA", "Contract Date",
     "Date contract was signed",
     14, None),
]

for col, label, instruction, width, num_fmt in pending_cols:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(ORANGE_BG),
          font=hdr_font(color="000000"),
          alignment=center(wrap=True))
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
    ws.column_dimensions[col].width = width
    if num_fmt:
        for row in range(5, 2001):
            ws[f"{col}{row}"].number_format = num_fmt

# ── CLAUDE OUTPUT (BB-BJ) ─────────────────────────────────────────────────────
# (col, label, instruction, width, num_fmt, formula, cell_fill)
claude_cols = [
    ("BB", "Comp Count",              "Auto — Make.com",        9,  None,     None, None),
    ("BC", "Starting Radius Used (mi)","Auto — Make.com",       13, None,     None, None),
    ("BD", "Final Radius Used (mi)",  "Auto — Make.com",        13, None,     None, None),
    ("BE", "Comp Age Window (days)",  "Auto — Make.com",        12, None,     None, None),
    ("BF", "Claude Price Low",        "Auto — Claude output",   13, "$#,##0", None, None),
    ("BG", "Claude Price High",       "Auto — Claude output",   13, "$#,##0", None, None),
    ("BH", "Claude Midpoint",         "Auto-calculated",        13, "$#,##0",
     '=IFERROR(IF(AND(BF5<>"",BG5<>""),AVERAGE(BF5,BG5),""),"—")',
     GREEN_BG),
    ("BI", "Confidence Level",        "Auto — Claude output",   13, None,     None, None),
    ("BJ", "Report Link",             "Auto — Google Drive link",24, None,    None, None),
]

for col, label, instruction, width, num_fmt, formula, cell_fill in claude_cols:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(NAVY_LIGHT),
          font=hdr_font(),
          alignment=center(wrap=True))
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
    ws.column_dimensions[col].width = width

    if num_fmt:
        for row in range(5, 2001):
            ws[f"{col}{row}"].number_format = num_fmt

    if formula:
        c = ws[f"{col}5"]
        c.value = formula
        c.number_format = "$#,##0"
        c.border = thin_border()
        if cell_fill:
            c.fill = solid(cell_fill)

# Confidence Level dropdown (BI)
dv = DataValidation(
    type="list",
    formula1='"HIGH,MEDIUM,LOW"',
    allow_blank=True
)
dv.sqref = "BI5:BI2000"
ws.add_data_validation(dv)

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 6 complete — Tab 1 cols 52-62 (AZ-BJ) built.")
print("  PENDING/CONTRACT banner AZ2:BA2 (ORANGE_BG)")
print("  CLAUDE OUTPUT banner BB2:BJ2 (NAVY_LIGHT)")
print("  Contract Price (AZ) formatted $#,##0.")
print("  Claude Price Low/High (BF/BG) formatted $#,##0.")
print("  Claude Midpoint (BH) formula seeded row 5, GREEN_BG, $#,##0.")
print("  Confidence Level (BI) dropdown: HIGH/MEDIUM/LOW.")
