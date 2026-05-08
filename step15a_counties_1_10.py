"""Step 15a: Counties Reference — states 1-10 (Alabama through Georgia)."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

NAVY  = "1A3A5C"
WHITE = "FFFFFF"

def solid(hex_color): return PatternFill("solid", fgColor=hex_color)
def thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

counties = {
    'Alabama': 'Autauga,Baldwin,Barbour,Bibb,Blount,Bullock,Butler,Calhoun,Chambers,Cherokee,Chilton,Choctaw,Clarke,Clay,Cleburne,Coffee,Colbert,Conecuh,Coosa,Covington,Crenshaw,Cullman,Dale,Dallas,DeKalb,Elmore,Escambia,Etowah,Fayette,Franklin,Geneva,Greene,Hale,Henry,Houston,Jackson,Jefferson,Lamar,Lauderdale,Lawrence,Lee,Limestone,Lowndes,Macon,Madison,Marengo,Marion,Marshall,Mobile,Monroe,Montgomery,Morgan,Perry,Pickens,Pike,Randolph,Russell,Shelby,St. Clair,Sumter,Talladega,Tallapoosa,Tuscaloosa,Walker,Washington,Wilcox,Winston',
    'Alaska': 'Aleutians East,Aleutians West,Anchorage,Bethel,Bristol Bay,Chugach,Copper River,Denali,Dillingham,Fairbanks North Star,Haines,Hoonah-Angoon,Juneau,Kenai Peninsula,Ketchikan Gateway,Kodiak Island,Kusilvak,Lake and Peninsula,Matanuska-Susitna,Nome,North Slope,Northwest Arctic,Petersburg,Prince of Wales-Hyder,Sitka,Skagway,Southeast Fairbanks,Wrangell,Yakutat,Yukon-Koyukuk',
    'Arizona': 'Apache,Cochise,Coconino,Gila,Graham,Greenlee,La Paz,Maricopa,Mohave,Navajo,Pima,Pinal,Santa Cruz,Yavapai,Yuma',
    'Arkansas': 'Arkansas,Ashley,Baxter,Benton,Boone,Bradley,Calhoun,Carroll,Chicot,Clark,Clay,Cleburne,Cleveland,Columbia,Conway,Craighead,Crawford,Crittenden,Cross,Dallas,Desha,Drew,Faulkner,Franklin,Fulton,Garland,Grant,Greene,Hempstead,Hot Spring,Howard,Independence,Izard,Jackson,Jefferson,Johnson,Lafayette,Lawrence,Lee,Lincoln,Little River,Logan,Lonoke,Madison,Marion,Miller,Mississippi,Monroe,Montgomery,Nevada,Newton,Ouachita,Perry,Phillips,Pike,Poinsett,Polk,Pope,Prairie,Pulaski,Randolph,St. Francis,Saline,Scott,Searcy,Sebastian,Sevier,Sharp,Stone,Union,Van Buren,Washington,White,Woodruff,Yell',
    'California': 'Alameda,Alpine,Amador,Butte,Calaveras,Colusa,Contra Costa,Del Norte,El Dorado,Fresno,Glenn,Humboldt,Imperial,Inyo,Kern,Kings,Lake,Lassen,Los Angeles,Madera,Marin,Mariposa,Mendocino,Merced,Modoc,Mono,Monterey,Napa,Nevada,Orange,Placer,Plumas,Riverside,Sacramento,San Benito,San Bernardino,San Diego,San Francisco,San Joaquin,San Luis Obispo,San Mateo,Santa Barbara,Santa Clara,Santa Cruz,Shasta,Sierra,Siskiyou,Solano,Sonoma,Stanislaus,Sutter,Tehama,Trinity,Tulare,Tuolumne,Ventura,Yolo,Yuba',
    'Colorado': 'Adams,Alamosa,Arapahoe,Archuleta,Baca,Bent,Boulder,Broomfield,Chaffee,Cheyenne,Clear Creek,Conejos,Costilla,Crowley,Custer,Delta,Denver,Dolores,Douglas,Eagle,El Paso,Elbert,Fremont,Garfield,Gilpin,Grand,Gunnison,Hinsdale,Huerfano,Jackson,Jefferson,Kiowa,Kit Carson,La Plata,Lake,Larimer,Las Animas,Lincoln,Logan,Mesa,Mineral,Moffat,Montezuma,Montrose,Morgan,Otero,Ouray,Park,Phillips,Pitkin,Prowers,Pueblo,Rio Blanco,Rio Grande,Routt,Saguache,San Juan,San Miguel,Sedgwick,Summit,Teller,Washington,Weld,Yuma',
    'Connecticut': 'Fairfield,Hartford,Litchfield,Middlesex,New Haven,New London,Tolland,Windham',
    'Delaware': 'Kent,New Castle,Sussex',
    'Florida': 'Alachua,Baker,Bay,Bradford,Brevard,Broward,Calhoun,Charlotte,Citrus,Clay,Collier,Columbia,DeSoto,Dixie,Duval,Escambia,Flagler,Franklin,Gadsden,Gilchrist,Glades,Gulf,Hamilton,Hardee,Hendry,Hernando,Highlands,Hillsborough,Holmes,Indian River,Jackson,Jefferson,Lafayette,Lake,Lee,Leon,Levy,Liberty,Madison,Manatee,Marion,Martin,Miami-Dade,Monroe,Nassau,Okaloosa,Okeechobee,Orange,Osceola,Palm Beach,Pasco,Pinellas,Polk,Putnam,Santa Rosa,Sarasota,Seminole,St. Johns,St. Lucie,Sumter,Suwannee,Taylor,Union,Volusia,Wakulla,Walton,Washington',
    'Georgia': 'Appling,Atkinson,Bacon,Baker,Baldwin,Banks,Barrow,Bartow,Ben Hill,Berrien,Bibb,Bleckley,Brantley,Brooks,Bryan,Bulloch,Burke,Butts,Calhoun,Camden,Candler,Carroll,Catoosa,Charlton,Chatham,Chattahoochee,Chattooga,Cherokee,Clarke,Clay,Clayton,Clinch,Cobb,Coffee,Colquitt,Columbia,Cook,Coweta,Crawford,Crisp,Dade,Dawson,Decatur,DeKalb,Dodge,Dooly,Dougherty,Douglas,Early,Echols,Effingham,Elbert,Emanuel,Evans,Fannin,Fayette,Floyd,Forsyth,Franklin,Fulton,Gilmer,Glascock,Glynn,Gordon,Grady,Greene,Gwinnett,Habersham,Hall,Hancock,Haralson,Harris,Hart,Heard,Henry,Houston,Irwin,Jackson,Jasper,Jeff Davis,Jefferson,Jenkins,Johnson,Jones,Lamar,Lanier,Laurens,Lee,Liberty,Lincoln,Long,Lowndes,Lumpkin,McDuffie,McIntosh,Macon,Madison,Marion,Meriwether,Miller,Mitchell,Monroe,Montgomery,Morgan,Murray,Muscogee,Newton,Oconee,Oglethorpe,Paulding,Peach,Pickens,Pierce,Pike,Polk,Pulaski,Putnam,Quitman,Rabun,Randolph,Richmond,Rockdale,Schley,Screven,Seminole,Spalding,Stephens,Stewart,Sumter,Talbot,Taliaferro,Tattnall,Taylor,Telfair,Terrell,Thomas,Tift,Toombs,Towns,Treutlen,Troup,Turner,Twiggs,Union,Upson,Walker,Walton,Ware,Warren,Washington,Wayne,Webster,Wheeler,White,Whitfield,Wilcox,Wilkes,Wilkinson,Worth',
}

wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Counties Reference"]
ws.sheet_state = "hidden"

# ── Title row 1 ───────────────────────────────────────────────────────────────
ws.merge_cells("A1:B1")
ws["A1"].value = "COUNTIES REFERENCE — DO NOT EDIT"
ws["A1"].fill  = solid(NAVY)
ws["A1"].font  = Font(bold=True, size=10, color=WHITE, name="Arial")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 20

# ── Headers row 2 ─────────────────────────────────────────────────────────────
ws["A2"].value = "State"
ws["B2"].value = "Counties"
for cell in [ws["A2"], ws["B2"]]:
    cell.font      = Font(bold=True, size=9, name="Arial")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border    = thin()

# ── Column widths ─────────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 300

# ── Write data starting row 3 ────────────────────────────────────────────────
start_row = 3
for i, (state, county_list) in enumerate(counties.items()):
    row = start_row + i
    ws[f"A{row}"].value = state
    ws[f"B{row}"].value = county_list
    ws[f"A{row}"].font  = Font(size=9, name="Arial")
    ws[f"B{row}"].font  = Font(size=9, name="Arial")
    ws[f"A{row}"].border = thin()
    ws[f"B{row}"].border = thin()

wb.save("SCG_CMA_System_v7.xlsx")

last_row = start_row + len(counties) - 1
print(f"Step 15a complete — {len(counties)} states written (rows 3-{last_row}).")
print(f"States written: {', '.join(counties.keys())}")
