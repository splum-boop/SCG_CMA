"""Step 12: Tab 5 — Accuracy Scorecard."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY        = "1A3A5C"
NAVY_LIGHT  = "2E5F8A"
PURPLE      = "4B2E83"
PURPLE_LIGHT= "E8E0F5"
LIGHT_GREY  = "F2F2F2"
MED_GREY    = "CCCCCC"
DARK_GREY   = "888888"
WHITE       = "FFFFFF"

PLACEHOLDER = "Populates after closed transactions"

# ── Helpers ───────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def apply(cell, value=None, fill=None, font=None, alignment=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    cell.border = thin_border()

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def hdr_font(color=WHITE, size=9, bold=True):
    return Font(bold=bold, size=size, color=color, name="Arial")

def placeholder_font():
    return Font(italic=True, size=8, color=DARK_GREY, name="Arial")

def banner(ws, start, end, text, bg, txt=WHITE, height=18):
    ws.merge_cells(f"{start}:{end}")
    apply(ws[start],
          value=text,
          fill=solid(bg),
          font=hdr_font(color=txt),
          alignment=left_wrap())
    row = int(''.join(filter(str.isdigit, start)))
    ws.row_dimensions[row].height = height

def write_headers(ws, row, cols, headers, bg, txt=WHITE):
    for col, label in zip(cols, headers):
        apply(ws[f"{col}{row}"],
              value=label,
              fill=solid(bg),
              font=hdr_font(color=txt),
              alignment=center(wrap=True))
        ws.column_dimensions[col].width = 22
    ws.row_dimensions[row].height = 36

def write_placeholder_row(ws, row, label_col, label, data_cols, row_bg):
    apply(ws[f"{label_col}{row}"],
          value=label,
          fill=solid(LIGHT_GREY),
          font=Font(bold=True, size=9, color="000000", name="Arial"),
          alignment=left_wrap())
    for col in data_cols:
        apply(ws[f"{col}{row}"],
              value=PLACEHOLDER,
              fill=solid(row_bg),
              font=placeholder_font(),
              alignment=center())
    ws.row_dimensions[row].height = 16

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Accuracy Scorecard"]

# 8 columns for sections 1-2-3, cols A-H
COLS8 = ["A","B","C","D","E","F","G","H"]
LAST  = "H"

# ── ROW 1 — Title Banner ──────────────────────────────────────────────────────
ws.merge_cells(f"A1:{LAST}1")
apply(ws["A1"],
      value="SCG CMA SYSTEM — ACCURACY SCORECARD | Auto-builds as transactions close | v7",
      fill=solid(NAVY),
      font=hdr_font(size=12),
      alignment=center())
ws.row_dimensions[1].height = 22

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — By Property Type
# ═══════════════════════════════════════════════════════════════════════════════
banner(ws, "A2", f"{LAST}2", "CLAUDE ACCURACY BY PROPERTY TYPE", NAVY_LIGHT)

s1_headers = ["Property Type","Reports Run","Avg Claude vs List %","Avg Claude vs Sold %",
              "Avg List vs Sold %","HIGH Conf Avg Var","MEDIUM Conf Avg Var","LOW Conf Avg Var"]
write_headers(ws, 3, COLS8, s1_headers, NAVY_LIGHT)

s1_types = ["In-Town Residential","Rural <2ac","Rural 2-10ac","Rural 10-40ac",
            "Rural 40-100ac","Rural 100+ac","Manufactured","Condo/Townhome",
            "Multi-Family","Farm/Land"]
for i, prop_type in enumerate(s1_types, start=4):
    write_placeholder_row(ws, i, "A", prop_type, ["B","C","D","E","F","G","H"], WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — By County
# ═══════════════════════════════════════════════════════════════════════════════
banner(ws, "A15", f"{LAST}15", "CLAUDE ACCURACY BY COUNTY", NAVY_LIGHT)

s2_headers = ["County","Reports Run","Avg Claude vs Sold %","Avg Bias Direction",
              "Best Type","Avg Comp Count","Avg Final Radius (mi)","Notes"]
write_headers(ws, 16, COLS8, s2_headers, NAVY_LIGHT)

s2_counties = ["Fayette","Marion","Shelby","Effingham (Rural)","Bond",
               "Montgomery","Clay","St. Clair (Metro East)","Madison (Metro East)",
               "[Add counties as needed]"]
for i, county in enumerate(s2_counties, start=17):
    write_placeholder_row(ws, i, "A", county, ["B","C","D","E","F","G","H"], WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — Bias Correction Summary
# ═══════════════════════════════════════════════════════════════════════════════
banner(ws, "A28", f"{LAST}28",
       "BIAS CORRECTION SUMMARY — Make.com feeds this into Claude before every report",
       PURPLE, height=18)

ws.merge_cells(f"A29:{LAST}29")
apply(ws["A29"],
      value=("Before each report Claude receives: Your last N reports in [Property Type] in [County] "
             "ran X% high/low on average. Adjust accordingly. "
             "Populates after 10+ closed transactions per category."),
      fill=solid(PURPLE_LIGHT),
      font=Font(italic=True, size=8, color=PURPLE, name="Arial"),
      alignment=left_wrap())
ws.row_dimensions[29].height = 30

s3_headers = ["Property Type + County","Sample Size","Avg Over %","Avg Under %",
              "Net Bias","Correction Applied","Last Updated","Notes"]
write_headers(ws, 30, COLS8, s3_headers, PURPLE)

s3_rows = [
    "In-Town Residential — Fayette County",
    "Rural 10-40ac — Fayette County",
    "Rural 40-100ac — Fayette County",
    "In-Town Residential — St. Clair County",
    "Rural 2-10ac — Marion County",
]
for i, label in enumerate(s3_rows, start=31):
    for col in COLS8:
        val = label if col == "A" else PLACEHOLDER
        apply(ws[f"{col}{i}"],
              value=val,
              fill=solid(PURPLE_LIGHT),
              font=(Font(size=9, color=PURPLE, name="Arial") if col == "A"
                    else placeholder_font()),
              alignment=left_wrap())
    ws.row_dimensions[i].height = 16

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — Overall KPIs
# ═══════════════════════════════════════════════════════════════════════════════
banner(ws, "A37", f"{LAST}37", "OVERALL KPIs", NAVY)

kpis = [
    ("Total Reports Generated",       "Count from Property Submissions"),
    ("Avg Claude vs Sold All",         PLACEHOLDER),
    ("Avg List-to-Sale In-Town",       PLACEHOLDER),
    ("Avg List-to-Sale Rural",         PLACEHOLDER),
    ("Avg DOM In-Town",                PLACEHOLDER),
    ("Avg DOM Rural",                  PLACEHOLDER),
    ("Prompt Version",                 "v7"),
    ("Last Updated",                   "2026-05-07"),
    ("Next Review",                    "After 20 closed transactions"),
]
for i, (label, value) in enumerate(kpis, start=38):
    apply(ws[f"A{i}"],
          value=label,
          fill=solid(NAVY_LIGHT if i % 2 == 0 else LIGHT_GREY),
          font=Font(bold=True, size=9, color=WHITE if i % 2 == 0 else "000000", name="Arial"),
          alignment=left_wrap())
    is_placeholder = value == PLACEHOLDER
    apply(ws[f"B{i}"],
          value=value,
          fill=solid(WHITE),
          font=placeholder_font() if is_placeholder else Font(size=9, color="000000", name="Arial"),
          alignment=left_wrap())
    # Fill remaining cols with border
    for col in ["C","D","E","F","G","H"]:
        ws[f"{col}{i}"].border = thin_border()
    ws.row_dimensions[i].height = 16

# ── Set column A width ────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 30

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 12 complete — Tab 5 Accuracy Scorecard built.")
print("  Section 1: Property Type rows 4-13")
print("  Section 2: County rows 17-26")
print("  Section 3: Bias Correction rows 31-35 (PURPLE_LIGHT)")
print("  Section 4: KPI label/value pairs rows 38-46")
