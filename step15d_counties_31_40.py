"""Step 15d: Counties Reference — states 31-40 (New Mexico through South Carolina)."""
import openpyxl
from openpyxl.styles import Font, Border, Side

def thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

counties = {
    'New Mexico': 'Bernalillo,Catron,Chaves,Cibola,Colfax,Curry,De Baca,Dona Ana,Eddy,Grant,Guadalupe,Harding,Hidalgo,Lea,Lincoln,Los Alamos,Luna,McKinley,Mora,Otero,Quay,Rio Arriba,Roosevelt,Sandoval,San Juan,San Miguel,Santa Fe,Sierra,Socorro,Taos,Torrance,Union,Valencia',
    'New York': 'Albany,Allegany,Bronx,Broome,Cattaraugus,Cayuga,Chautauqua,Chemung,Chenango,Clinton,Columbia,Cortland,Delaware,Dutchess,Erie,Essex,Franklin,Fulton,Genesee,Greene,Hamilton,Herkimer,Jefferson,Kings,Lewis,Livingston,Madison,Monroe,Montgomery,Nassau,New York,Niagara,Oneida,Onondaga,Ontario,Orange,Orleans,Oswego,Otsego,Putnam,Queens,Rensselaer,Richmond,Rockland,St. Lawrence,Saratoga,Schenectady,Schoharie,Schuyler,Seneca,Steuben,Suffolk,Sullivan,Tioga,Tompkins,Ulster,Warren,Washington,Wayne,Westchester,Wyoming,Yates',
    'North Carolina': 'Alamance,Alexander,Alleghany,Anson,Ashe,Avery,Beaufort,Bertie,Bladen,Brunswick,Buncombe,Burke,Cabarrus,Caldwell,Camden,Carteret,Caswell,Catawba,Chatham,Cherokee,Chowan,Clay,Cleveland,Columbus,Craven,Cumberland,Currituck,Dare,Davidson,Davie,Duplin,Durham,Edgecombe,Forsyth,Franklin,Gaston,Gates,Graham,Granville,Greene,Guilford,Halifax,Harnett,Haywood,Henderson,Hertford,Hoke,Hyde,Iredell,Jackson,Johnston,Jones,Lee,Lenoir,Lincoln,Macon,Madison,Martin,McDowell,Mecklenburg,Mitchell,Montgomery,Moore,Nash,New Hanover,Northampton,Onslow,Orange,Pamlico,Pasquotank,Pender,Perquimans,Person,Pitt,Polk,Randolph,Richmond,Robeson,Rockingham,Rowan,Rutherford,Sampson,Scotland,Stanly,Stokes,Surry,Swain,Transylvania,Tyrrell,Union,Vance,Wake,Warren,Washington,Watauga,Wayne,Wilkes,Wilson,Yadkin,Yancey',
    'North Dakota': 'Adams,Barnes,Benson,Billings,Bottineau,Bowman,Burke,Burleigh,Cass,Cavalier,Dickey,Divide,Dunn,Eddy,Emmons,Foster,Golden Valley,Grand Forks,Grant,Griggs,Hettinger,Kidder,LaMoure,Logan,McHenry,McIntosh,McKenzie,McLean,Mercer,Morton,Mountrail,Nelson,Oliver,Pembina,Pierce,Ramsey,Ransom,Renville,Richland,Rolette,Sargent,Sheridan,Sioux,Slope,Stark,Steele,Stutsman,Towner,Traill,Walsh,Ward,Wells,Williams',
    'Ohio': 'Adams,Allen,Ashland,Ashtabula,Athens,Auglaize,Belmont,Brown,Butler,Carroll,Champaign,Clark,Clermont,Clinton,Columbiana,Coshocton,Crawford,Cuyahoga,Darke,Defiance,Delaware,Erie,Fairfield,Fayette,Franklin,Fulton,Gallia,Geauga,Greene,Guernsey,Hamilton,Hancock,Hardin,Harrison,Henry,Highland,Hocking,Holmes,Huron,Jackson,Jefferson,Knox,Lake,Lawrence,Licking,Logan,Lorain,Lucas,Madison,Mahoning,Marion,Medina,Meigs,Mercer,Miami,Monroe,Montgomery,Morgan,Morrow,Muskingum,Noble,Ottawa,Paulding,Perry,Pickaway,Pike,Portage,Preble,Putnam,Richland,Ross,Sandusky,Scioto,Seneca,Shelby,Stark,Summit,Trumbull,Tuscarawas,Union,Van Wert,Vinton,Warren,Washington,Wayne,Williams,Wood,Wyandot',
    'Oklahoma': 'Adair,Alfalfa,Atoka,Beaver,Beckham,Blaine,Bryan,Caddo,Canadian,Carter,Cherokee,Choctaw,Cimarron,Cleveland,Coal,Comanche,Cotton,Craig,Creek,Custer,Delaware,Dewey,Ellis,Garfield,Garvin,Grady,Grant,Greer,Harmon,Harper,Haskell,Hughes,Jackson,Jefferson,Johnston,Kay,Kingfisher,Kiowa,Latimer,Le Flore,Lincoln,Logan,Love,Major,Marshall,Mayes,McClain,McCurtain,McIntosh,Murray,Muskogee,Noble,Nowata,Okfuskee,Oklahoma,Okmulgee,Osage,Ottawa,Pawnee,Payne,Pittsburg,Pontotoc,Pottawatomie,Pushmataha,Roger Mills,Rogers,Seminole,Sequoyah,Stephens,Texas,Tillman,Tulsa,Wagoner,Washington,Washita,Woods,Woodward',
    'Oregon': 'Baker,Benton,Clackamas,Clatsop,Columbia,Coos,Crook,Curry,Deschutes,Douglas,Gilliam,Grant,Harney,Hood River,Jackson,Jefferson,Josephine,Klamath,Lake,Lane,Lincoln,Linn,Malheur,Marion,Morrow,Multnomah,Polk,Sherman,Tillamook,Umatilla,Union,Wallowa,Wasco,Washington,Wheeler,Yamhill',
    'Pennsylvania': 'Adams,Allegheny,Armstrong,Beaver,Bedford,Berks,Blair,Bradford,Bucks,Butler,Cambria,Cameron,Carbon,Centre,Chester,Clarion,Clearfield,Clinton,Columbia,Crawford,Cumberland,Dauphin,Delaware,Elk,Erie,Fayette,Forest,Franklin,Fulton,Greene,Huntingdon,Indiana,Jefferson,Juniata,Lackawanna,Lancaster,Lawrence,Lebanon,Lehigh,Luzerne,Lycoming,McKean,Mercer,Mifflin,Monroe,Montgomery,Montour,Northampton,Northumberland,Perry,Philadelphia,Pike,Potter,Schuylkill,Snyder,Somerset,Sullivan,Susquehanna,Tioga,Union,Venango,Warren,Washington,Wayne,Westmoreland,Wyoming,York',
    'Rhode Island': 'Bristol,Kent,Newport,Providence,Washington',
    'South Carolina': 'Abbeville,Aiken,Allendale,Anderson,Bamberg,Barnwell,Beaufort,Berkeley,Calhoun,Charleston,Cherokee,Chester,Chesterfield,Clarendon,Colleton,Darlington,Dillon,Dorchester,Edgefield,Fairfield,Florence,Georgetown,Greenville,Greenwood,Hampton,Horry,Jasper,Kershaw,Lancaster,Laurens,Lee,Lexington,Marion,Marlboro,McCormick,Newberry,Oconee,Orangeburg,Pickens,Richland,Saluda,Spartanburg,Sumter,Union,Williamsburg,York',
}

wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Counties Reference"]
ws.sheet_state = "hidden"

next_row = 3
while ws[f"A{next_row}"].value:
    next_row += 1

for state, county_list in counties.items():
    ws[f"A{next_row}"].value  = state
    ws[f"B{next_row}"].value  = county_list
    ws[f"A{next_row}"].font   = Font(size=9, name="Arial")
    ws[f"B{next_row}"].font   = Font(size=9, name="Arial")
    ws[f"A{next_row}"].border = thin()
    ws[f"B{next_row}"].border = thin()
    next_row += 1

wb.save("SCG_CMA_System_v7.xlsx")
last = next_row - 1
print(f"Step 15d complete — {len(counties)} states appended (rows 33-{last}).")
print(f"States: {', '.join(counties.keys())}")
print(f"Total data rows so far: {last - 2}")
