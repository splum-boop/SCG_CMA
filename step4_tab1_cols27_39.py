"""Step 4: Tab 1 — cols 27-39 (AA-AM), Structure Features + Lot/Acreage."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule, FormulaRule

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY_LIGHT  = "2E5F8A"
GOLD        = "C9A84C"
INSTRUCT_BG = "EAF0F6"
MED_GREY    = "CCCCCC"
DARK_GREY   = "888888"
WHITE       = "FFFFFF"
GREEN_BG    = "E2EFDA"
SKIP_GREY   = "D9D9D9"

# ── Helpers ───────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def apply(cell, value=None, fill=None, font=None, alignment=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    cell.border = thin_border()

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def hdr_font(color=WHITE):
    return Font(bold=True, size=9, color=color, name="Arial")

def instr_font():
    return Font(italic=True, size=7, color=DARK_GREY, name="Arial")

# ── Column map ────────────────────────────────────────────────────────────────
# Col 27=AA, 28=AB, 29=AC, 30=AD, 31=AE, 32=AF
# Col 33=AG, 34=AH, 35=AI, 36=AJ, 37=AK, 38=AL, 39=AM

# (col_letter, header_label, instruction, header_bg, width, dropdown_formula)
structure_cols = [
    ("AA", "Basement",      "Select Yes or No",            NAVY_LIGHT, 10,
     '"Yes,No"'),
    ("AB", "Basement Type", "Select if Basement = Yes",    NAVY_LIGHT, 16,
     '"Unfinished,Partially Finished,Fully Finished,Partial Basement,Cellar"'),
    ("AC", "Garage",        "Select garage type",          NAVY_LIGHT, 16,
     '"None,1 Car Attached,2 Car Attached,3+ Car Attached,Detached,Multiple"'),
    ("AD", "Stories",       "Select number of stories",    NAVY_LIGHT, 13,
     '"Single,Multi,Split Level"'),
    ("AE", "Central Air",   "Yes or No",                   NAVY_LIGHT, 10,
     '"Yes,No"'),
    ("AF", "Heat Type",     "Select heat source",          NAVY_LIGHT, 14,
     '"Natural Gas,Electric,Propane,Other"'),
]

lot_cols = [
    # (col, header, instruction, header_bg, width, formula_or_None, cell_fill)
    ("AG", "Lot Length (ft)",    "OPTION 1 — Enter length in feet only",       GOLD,       13, None,    None),
    ("AH", "Lot Width (ft)",     "OPTION 1 — Enter width in feet only",        GOLD,       13, None,    None),
    ("AI", "Lot Sq Ft (direct)", "OPTION 2 — Enter sq ft if known",            GOLD,       13, None,    None),
    ("AJ", "Lot Acreage (direct)","OPTION 3 — Enter acreage if known",         GOLD,       13, None,    None),
    ("AK", "CALC: Lot Sq Ft",   "Auto-calculated — do not edit",              NAVY_LIGHT, 13,
     "=IFERROR(IF(AND(AG5<>\"\",AH5<>\"\"),AG5*AH5,IF(AI5<>\"\",AI5,IF(AJ5<>\"\",AJ5*43560,\"\"))),"
     "\"ERROR\")",
     GREEN_BG),
    ("AL", "CALC: Lot Acres",   "Auto-calculated — do not edit",              NAVY_LIGHT, 13,
     '=IFERROR(IF(AK5<>"",AK5/43560,""),"—")',
     GREEN_BG),
    ("AM", "Lot Notes",         "Corner lot, alley, oversized, etc.",          GOLD,       22, None,    None),
]

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Property Submissions"]

# ── ROW 2 — Section Banners ───────────────────────────────────────────────────
ws.merge_cells("AA2:AF2")
apply(ws["AA2"],
      value="STRUCTURE FEATURES",
      fill=solid(NAVY_LIGHT),
      font=hdr_font(),
      alignment=center())

ws.merge_cells("AG2:AM2")
apply(ws["AG2"],
      value="LOT / ACREAGE — Fill ONE option only: cols AG+AH, AI, or AJ",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center())

# ── ROW 3 — Headers + ROW 4 — Instructions (Structure Features) ───────────────
for col, label, instruction, bg, _, dv_formula in structure_cols:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(bg),
          font=hdr_font(),
          alignment=center(wrap=True))
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
    ws.column_dimensions[col].width = _
    # Dropdown
    dv = DataValidation(type="list", formula1=dv_formula, allow_blank=True)
    dv.sqref = f"{col}5:{col}2000"
    ws.add_data_validation(dv)

# ── ROW 3 — Headers + ROW 4 — Instructions (Lot/Acreage) ─────────────────────
for col, label, instruction, bg, width, formula, cell_fill in lot_cols:
    # Header
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(bg),
          font=hdr_font(color="000000" if bg == GOLD else WHITE),
          alignment=center(wrap=True))
    # Instruction
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
    ws.column_dimensions[col].width = width

    # Seed formula in row 5 data start cell
    if formula:
        c = ws[f"{col}5"]
        c.value = formula
        if cell_fill:
            c.fill = solid(cell_fill)
        c.border = thin_border()
        if col == "AL":
            c.number_format = "0.000"

# ── Conditional Formatting — Basement Type (AB) grey when Basement (AA) = No ──
# Formula: =$AA5="No"  →  applied to AB5:AB2000
grey_fill = solid(SKIP_GREY)
grey_font = Font(italic=True, color=DARK_GREY, name="Arial")

from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add(
    "AB5:AB2000",
    FormulaRule(formula=['$AA5="No"'], fill=grey_fill, font=grey_font)
)

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 4 complete — Tab 1 cols 27-39 (AA-AM) built.")
print("  STRUCTURE FEATURES banner AA2:AF2 (NAVY_LIGHT)")
print("  LOT / ACREAGE banner AG2:AM2 (GOLD)")
print("  6 dropdowns added for structure cols.")
print("  CALC formulas seeded in AK5 and AL5.")
print("  Conditional formatting: AB grey when AA = No.")
