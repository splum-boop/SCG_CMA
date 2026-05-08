"""Step 15c: Counties Reference — states 21-30 (Massachusetts through New Jersey)."""
import openpyxl
from openpyxl.styles import Font, Border, Side

def thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

counties = {
    'Massachusetts': 'Barnstable,Berkshire,Bristol,Dukes,Essex,Franklin,Hampden,Hampshire,Middlesex,Nantucket,Norfolk,Plymouth,Suffolk,Worcester',
    'Michigan': 'Alcona,Alger,Allegan,Alpena,Antrim,Arenac,Baraga,Barry,Bay,Benzie,Berrien,Branch,Calhoun,Cass,Charlevoix,Cheboygan,Chippewa,Clare,Clinton,Crawford,Delta,Dickinson,Eaton,Emmet,Genesee,Gladwin,Gogebic,Grand Traverse,Gratiot,Hillsdale,Houghton,Huron,Ingham,Ionia,Iosco,Iron,Isabella,Jackson,Kalamazoo,Kalkaska,Kent,Keweenaw,Lake,Lapeer,Leelanau,Lenawee,Livingston,Luce,Mackinac,Macomb,Manistee,Marquette,Mason,Mecosta,Menominee,Midland,Missaukee,Monroe,Montcalm,Montmorency,Muskegon,Newaygo,Oakland,Oceana,Ogemaw,Ontonagon,Osceola,Oscoda,Otsego,Ottawa,Presque Isle,Roscommon,Saginaw,Sanilac,Schoolcraft,Shiawassee,St. Clair,St. Joseph,Tuscola,Van Buren,Washtenaw,Wayne,Wexford',
    'Minnesota': 'Aitkin,Anoka,Becker,Beltrami,Benton,Big Stone,Blue Earth,Brown,Carlton,Carver,Cass,Chippewa,Chisago,Clay,Clearwater,Cook,Cottonwood,Crow Wing,Dakota,Dodge,Douglas,Faribault,Fillmore,Freeborn,Goodhue,Grant,Hennepin,Houston,Hubbard,Isanti,Itasca,Jackson,Kanabec,Kandiyohi,Kittson,Koochiching,Lac qui Parle,Lake,Lake of the Woods,Le Sueur,Lincoln,Lyon,Mahnomen,Marshall,Martin,McLeod,Meeker,Mille Lacs,Morrison,Mower,Murray,Nicollet,Nobles,Norman,Olmsted,Otter Tail,Pennington,Pine,Pipestone,Polk,Pope,Ramsey,Red Lake,Redwood,Renville,Rice,Rock,Roseau,Scott,Sherburne,Sibley,St. Louis,Stearns,Steele,Stevens,Swift,Todd,Traverse,Wabasha,Wadena,Waseca,Washington,Watonwan,Wilkin,Winona,Wright,Yellow Medicine',
    'Mississippi': 'Adams,Alcorn,Amite,Attala,Benton,Bolivar,Calhoun,Carroll,Chickasaw,Choctaw,Claiborne,Clarke,Clay,Coahoma,Copiah,Covington,De Soto,Forrest,Franklin,George,Greene,Grenada,Hancock,Harrison,Hinds,Holmes,Humphreys,Issaquena,Itawamba,Jackson,Jasper,Jefferson,Jefferson Davis,Jones,Kemper,Lafayette,Lamar,Lauderdale,Lawrence,Leake,Lee,Leflore,Lincoln,Lowndes,Madison,Marion,Marshall,Monroe,Montgomery,Neshoba,Newton,Noxubee,Oktibbeha,Panola,Pearl River,Perry,Pike,Pontotoc,Prentiss,Quitman,Rankin,Scott,Sharkey,Simpson,Smith,Stone,Sunflower,Tallahatchie,Tate,Tippah,Tishomingo,Tunica,Union,Walthall,Warren,Washington,Wayne,Webster,Wilkinson,Winston,Yalobusha,Yazoo',
    'Missouri': 'Adair,Andrew,Atchison,Audrain,Barry,Barton,Bates,Benton,Bollinger,Boone,Buchanan,Butler,Caldwell,Callaway,Camden,Cape Girardeau,Carroll,Carter,Cass,Cedar,Chariton,Christian,Clark,Clay,Clinton,Cole,Cooper,Crawford,Dade,Dallas,Daviess,DeKalb,Dent,Douglas,Dunklin,Franklin,Gasconade,Gentry,Greene,Grundy,Harrison,Henry,Hickory,Holt,Howard,Howell,Iron,Jackson,Jasper,Jefferson,Johnson,Knox,Laclede,Lafayette,Lawrence,Lewis,Lincoln,Linn,Livingston,Macon,Madison,Maries,Marion,McDonald,Mercer,Miller,Mississippi,Moniteau,Monroe,Montgomery,Morgan,New Madrid,Newton,Nodaway,Oregon,Osage,Ozark,Pemiscot,Perry,Pettis,Phelps,Pike,Platte,Polk,Pulaski,Putnam,Ralls,Randolph,Ray,Reynolds,Ripley,St. Charles,St. Clair,St. Francois,St. Louis,St. Louis City,Ste. Genevieve,Saline,Schuyler,Scotland,Scott,Shannon,Shelby,Stoddard,Stone,Sullivan,Taney,Texas,Vernon,Warren,Washington,Wayne,Webster,Worth,Wright',
    'Montana': 'Beaverhead,Big Horn,Blaine,Broadwater,Carbon,Carter,Cascade,Chouteau,Custer,Daniels,Dawson,Deer Lodge,Fallon,Fergus,Flathead,Gallatin,Garfield,Glacier,Golden Valley,Granite,Hill,Jefferson,Judith Basin,Lake,Lewis and Clark,Liberty,Lincoln,Madison,McCone,Meagher,Mineral,Missoula,Musselshell,Park,Petroleum,Phillips,Pondera,Powder River,Powell,Prairie,Ravalli,Richland,Roosevelt,Rosebud,Sanders,Sheridan,Silver Bow,Stillwater,Sweet Grass,Teton,Toole,Treasure,Valley,Wheatland,Wibaux,Yellowstone',
    'Nebraska': 'Adams,Antelope,Arthur,Banner,Blaine,Boone,Box Butte,Boyd,Brown,Buffalo,Burt,Butler,Cass,Cedar,Chase,Cherry,Cheyenne,Clay,Colfax,Cuming,Custer,Dakota,Dawes,Dawson,Deuel,Dixon,Dodge,Douglas,Dundy,Fillmore,Franklin,Frontier,Furnas,Gage,Garden,Garfield,Gosper,Grant,Greeley,Hall,Hamilton,Harlan,Hayes,Hitchcock,Holt,Hooker,Howard,Jefferson,Johnson,Kearney,Keith,Keya Paha,Kimball,Knox,Lancaster,Lincoln,Logan,Loup,Madison,McPherson,Merrick,Morrill,Nance,Nemaha,Nuckolls,Otoe,Pawnee,Perkins,Phelps,Pierce,Platte,Polk,Red Willow,Richardson,Rock,Saline,Sarpy,Saunders,Scotts Bluff,Seward,Sheridan,Sherman,Sioux,Stanton,Thayer,Thomas,Thurston,Valley,Washington,Wayne,Webster,Wheeler,York',
    'Nevada': 'Carson City,Churchill,Clark,Douglas,Elko,Esmeralda,Eureka,Humboldt,Lander,Lincoln,Lyon,Mineral,Nye,Pershing,Storey,Washoe,White Pine',
    'New Hampshire': 'Belknap,Carroll,Cheshire,Coos,Grafton,Hillsborough,Merrimack,Rockingham,Strafford,Sullivan',
    'New Jersey': 'Atlantic,Bergen,Burlington,Camden,Cape May,Cumberland,Essex,Gloucester,Hudson,Hunterdon,Mercer,Middlesex,Monmouth,Morris,Ocean,Passaic,Salem,Somerset,Sussex,Union,Warren',
}

wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Counties Reference"]
ws.sheet_state = "hidden"

# Find next empty row
next_row = 3
while ws[f"A{next_row}"].value:
    next_row += 1

for state, county_list in counties.items():
    ws[f"A{next_row}"].value  = state
    ws[f"B{next_row}"].value  = county_list
    ws[f"A{next_row}"].font   = Font(size=9, name="Arial")
    ws[f"B{next_row}"].font   = Font(size=9, name="Arial")
    ws[f"A{next_row}"].border = thin()
    ws[f"B{next_row}"].border = thin()
    next_row += 1

wb.save("SCG_CMA_System_v7.xlsx")
last = next_row - 1
print(f"Step 15c complete — {len(counties)} states appended (rows 23-{last}).")
print(f"States: {', '.join(counties.keys())}")
print(f"Total data rows so far: {last - 2}")
