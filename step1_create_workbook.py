"""Step 1: Create workbook with 8 named tabs, tab 8 hidden."""
import openpyxl

wb = openpyxl.Workbook()

tab_names = [
    "Property Submissions",
    "Agent Roster",
    "Office Profiles",
    "Comp Library",
    "Accuracy Scorecard",
    "Velocity Metrics",
    "Monitoring Queue",
    "Counties Reference",
]

# Rename the default sheet to the first tab
wb.active.title = tab_names[0]

# Create remaining tabs
for name in tab_names[1:]:
    wb.create_sheet(title=name)

# Hide tab 8
wb["Counties Reference"].sheet_state = "hidden"

# Save
out = "/home/user/SCG_CMA/SCG_CMA_System_v7.xlsx"
wb.save(out)
print(f"Saved: {out}")
print(f"Sheets: {wb.sheetnames}")
print(f"Counties Reference state: {wb['Counties Reference'].sheet_state}")
