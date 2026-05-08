"""Step 2: Tab 1 — Property Submissions, cols 1-16, title/banners/headers/instructions."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── Colors ──────────────────────────────────────────────────────────────────
NAVY       = "1A3A5C"
NAVY_LIGHT = "2E5F8A"
INSTRUCT_BG = "EAF0F6"
MED_GREY   = "CCCCCC"
DARK_GREY  = "888888"
WHITE      = "FFFFFF"

TOTAL_COLS = 74  # full width for title merge

# ── Helpers ──────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_font(bold=True, size=9, color=WHITE):
    return Font(bold=bold, size=size, color=color, name="Arial")

def instr_font():
    return Font(italic=True, size=7, color=DARK_GREY, name="Arial")

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def apply(cell, value=None, fill=None, font=None, alignment=None, border=True):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border:
        cell.border = thin_border()

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Property Submissions"]

last_col = get_column_letter(TOTAL_COLS)   # BV

# ── ROW 1 — Title Banner ──────────────────────────────────────────────────────
ws.merge_cells(f"A1:{last_col}1")
apply(ws["A1"],
      value="SCG CMA SYSTEM — PROPERTY SUBMISSIONS | Agent-Facing | v7",
      fill=solid(NAVY),
      font=Font(bold=True, size=12, color=WHITE, name="Arial"),
      alignment=center())
ws.row_dimensions[1].height = 22

# ── ROW 2 — Section Banners (cols 1-16 only for now) ─────────────────────────
sections = [
    ("A2", "D2", "SUBMISSION INFO",       NAVY),
    ("E2", "P2", "PROPERTY DETAILS CORE", NAVY_LIGHT),
]
for start, end, label, color in sections:
    ws.merge_cells(f"{start}:{end}")
    apply(ws[start],
          value=label,
          fill=solid(color),
          font=hdr_font(bold=True, size=9),
          alignment=center())
ws.row_dimensions[2].height = 18

# ── ROW 3 — Column Headers ────────────────────────────────────────────────────
headers = [
    #  col   label               bg_color
    ("A", "Submission ID",    NAVY),
    ("B", "Submission Date",  NAVY),
    ("C", "Agent Name",       NAVY),
    ("D", "Agent Email",      NAVY),
    ("E", "Property Address", NAVY_LIGHT),
    ("F", "City",             NAVY_LIGHT),
    ("G", "County",           NAVY_LIGHT),
    ("H", "Zip Code",         NAVY_LIGHT),
    ("I", "Property Class",   NAVY_LIGHT),
    ("J", "Property Type",    NAVY_LIGHT),
    ("K", "Beds",             NAVY_LIGHT),
    ("L", "Baths",            NAVY_LIGHT),
    ("M", "Half Baths",       NAVY_LIGHT),
    ("N", "House Sq Ft",      NAVY_LIGHT),
    ("O", "Year Built",       NAVY_LIGHT),
    ("P", "Condition",        NAVY_LIGHT),
]
for col, label, bg in headers:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(bg),
          font=hdr_font(),
          alignment=center(wrap=True))
ws.row_dimensions[3].height = 44

# ── ROW 4 — Instructions ──────────────────────────────────────────────────────
instructions = [
    ("A", "Auto — Make.com"),
    ("B", "Auto — Make.com"),
    ("C", "Your full name — must match Agent Roster exactly"),
    ("D", "Your email address"),
    ("E", "Full street address"),
    ("F", "City name"),
    ("G", "County name"),
    ("H", "5-digit zip"),
    ("I", "Select property class"),
    ("J", "Select from dropdown"),
    ("K", "# bedrooms"),
    ("L", "Full baths"),
    ("M", "Half baths (0 if none)"),
    ("N", "Living area sq ft from MLS"),
    ("O", "4-digit year — original construction"),
    ("P", "Select from dropdown"),
]
for col, text in instructions:
    apply(ws[f"{col}4"],
          value=text,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
ws.row_dimensions[4].height = 30

# ── Dropdowns ─────────────────────────────────────────────────────────────────
dvs = [
    ("I5:I2000", '"In-Town Residential,Rural/Acreage"'),
    ("J5:J2000", '"Single Family - Stick Built,Single Family - Manufactured,Condo/Townhome,Multi-Family,Farm,Land"'),
    ("P5:P2000", '"Excellent,Good,Fair,Poor"'),
]
for sqref, formula in dvs:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.sqref = sqref
    ws.add_data_validation(dv)

# ── Column Widths ─────────────────────────────────────────────────────────────
widths = {
    "A": 13, "B": 14, "C": 18, "D": 22,
    "E": 30, "F": 16, "G": 14, "H": 10,
    "I": 16, "J": 20, "K":  7, "L":  7,
    "M":  8, "N": 12, "O": 11, "P": 12,
}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ── Freeze Panes ──────────────────────────────────────────────────────────────
ws.freeze_panes = "A5"

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 2 complete — Tab 1 cols 1-16 built.")
print(f"  Rows 1-4 set, freeze at A5, {len(dvs)} dropdowns added.")
print(f"  Title banner spans A1:{last_col}1 ({TOTAL_COLS} columns).")
