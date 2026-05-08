"""Step 5: Tab 1 — cols 40-51 (AN-AY), Rural/Acreage Fields + Comp Search Parameters."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY        = "1A3A5C"
GOLD        = "C9A84C"
INSTRUCT_BG = "EAF0F6"
MED_GREY    = "CCCCCC"
DARK_GREY   = "888888"
WHITE       = "FFFFFF"
GREEN_BG    = "E2EFDA"
YELLOW_BG   = "FFF2CC"
SKIP_GREY   = "D9D9D9"
LIGHT_BLUE  = "D5E8F4"

# ── Helpers ───────────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=MED_GREY)
    return Border(left=s, right=s, top=s, bottom=s)

def apply(cell, value=None, fill=None, font=None, alignment=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    cell.border = thin_border()

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def hdr_font(color=WHITE):
    return Font(bold=True, size=9, color=color, name="Arial")

def instr_font():
    return Font(italic=True, size=7, color=DARK_GREY, name="Arial")

def add_dv(ws, col, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.sqref = f"{col}5:{col}2000"
    ws.add_data_validation(dv)

# ── Load workbook ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("SCG_CMA_System_v7.xlsx")
ws = wb["Property Submissions"]

# ── Column map ────────────────────────────────────────────────────────────────
# Col 40=AN, 41=AO, 42=AP, 43=AQ, 44=AR
# Col 45=AS, 46=AT, 47=AU, 48=AV, 49=AW, 50=AX, 51=AY

# ── ROW 2 — Section Banners ───────────────────────────────────────────────────
ws.merge_cells("AN2:AR2")
apply(ws["AN2"],
      value="RURAL / ACREAGE FIELDS — Skip if In-Town Residential",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center())

ws.merge_cells("AS2:AY2")
apply(ws["AS2"],
      value="COMP SEARCH PARAMETERS",
      fill=solid(NAVY),
      font=hdr_font(),
      alignment=center())

# ── RURAL / ACREAGE FIELDS (AN-AR) ───────────────────────────────────────────

# AN — Acreage Band (formula, GREEN_BG, GOLD header)
apply(ws["AN3"],
      value="Acreage Band",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center(wrap=True))
apply(ws["AN4"],
      value="Auto-assigned — do not edit",
      fill=solid(INSTRUCT_BG),
      font=instr_font(),
      alignment=left_wrap())
ws.column_dimensions["AN"].width = 16
# Seed formula in row 5
c = ws["AN5"]
c.value = ('=IFERROR(IF(AL5="","",IF(AL5<2,"Residential (<2ac)",'
           'IF(AL5<10,"Small Rural (2-10ac)",IF(AL5<40,"Rural (10-40ac)",'
           'IF(AL5<100,"Farm/Estate (40-100ac)","Large Farm (100+ac)"))))),"—")')
c.fill  = solid(GREEN_BG)
c.border = thin_border()

# AO — Land Type (dropdown)
apply(ws["AO3"],
      value="Land Type",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center(wrap=True))
apply(ws["AO4"],
      value="RURAL ONLY",
      fill=solid(INSTRUCT_BG),
      font=instr_font(),
      alignment=left_wrap())
ws.column_dimensions["AO"].width = 16
add_dv(ws, "AO", '"Tillable Farmland,Timber,Pasture,Recreational,Mixed,Residential Lot"')

# AP — Primary Value Driver (dropdown)
apply(ws["AP3"],
      value="Primary Value Driver",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center(wrap=True))
apply(ws["AP4"],
      value="RURAL ONLY",
      fill=solid(INSTRUCT_BG),
      font=instr_font(),
      alignment=left_wrap())
ws.column_dimensions["AP"].width = 14
add_dv(ws, "AP", '"Land,Structure,Equal"')

# AQ — Outbuildings (free text)
apply(ws["AQ3"],
      value="Outbuildings",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center(wrap=True))
apply(ws["AQ4"],
      value="RURAL ONLY — Describe type and size",
      fill=solid(INSTRUCT_BG),
      font=instr_font(),
      alignment=left_wrap())
ws.column_dimensions["AQ"].width = 22

# AR — Water Features (dropdown)
apply(ws["AR3"],
      value="Water Features",
      fill=solid(GOLD),
      font=hdr_font(color="000000"),
      alignment=center(wrap=True))
apply(ws["AR4"],
      value="RURAL ONLY",
      fill=solid(INSTRUCT_BG),
      font=instr_font(),
      alignment=left_wrap())
ws.column_dimensions["AR"].width = 14
add_dv(ws, "AR", '"None,Pond,Creek,River Frontage,Multiple"')

# ── COMP SEARCH PARAMETERS (AS-AY) ───────────────────────────────────────────
comp_cols = [
    ("AS", "State",
     "Auto-populated from Office Profile — do not edit",
     NAVY, 12, None, LIGHT_BLUE),
    ("AT", "Primary Comp County",
     "Select county for comp search — required",
     NAVY, 20, None, None),
    ("AU", "Comp County 2",
     "Optional — add if property near county border",
     NAVY, 20, None, None),
    ("AV", "Comp County 3",
     "Optional — add third county if needed (e.g. Pana IL near 3 county junction)",
     NAVY, 20, None, None),
    ("AW", "Requested Starting Radius",
     "Override system default radius if needed",
     NAVY, 18,
     '"1 mile,2 miles,5 miles,10 miles,15 miles,20 miles,35 miles,50 miles,Let System Decide"',
     None),
    ("AX", "Additional City Exclusions",
     "Cities beyond office standard list to exclude",
     NAVY, 24, None, None),
    ("AY", "Agent Local Knowledge Notes",
     "Anything comps won't show — gravel road, best view, new well, fiber, rebuild cost, unique features",
     NAVY, 36, None, None),
]

for col, label, instruction, bg, width, dv_formula, data_fill in comp_cols:
    apply(ws[f"{col}3"],
          value=label,
          fill=solid(bg),
          font=hdr_font(),
          alignment=center(wrap=True))
    apply(ws[f"{col}4"],
          value=instruction,
          fill=solid(INSTRUCT_BG),
          font=instr_font(),
          alignment=left_wrap())
    ws.column_dimensions[col].width = width
    if dv_formula:
        add_dv(ws, col, dv_formula)

# State col (AS) — LIGHT_BLUE always on data rows (formula =TRUE)
ws.conditional_formatting.add(
    "AS5:AS2000",
    FormulaRule(formula=["TRUE"], fill=solid(LIGHT_BLUE))
)

# ── Conditional Formatting — Rural cols AN:AR ─────────────────────────────────
# In-Town Residential → SKIP_GREY + italic dark grey
# Rural/Acreage       → YELLOW_BG
grey_fill   = solid(SKIP_GREY)
grey_font   = Font(italic=True, color=DARK_GREY, name="Arial")
yellow_fill = solid(YELLOW_BG)

for rng in ["AN5:AN2000", "AO5:AO2000", "AP5:AP2000", "AQ5:AQ2000", "AR5:AR2000"]:
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=['$I5="In-Town Residential"'],
                    fill=grey_fill, font=grey_font)
    )
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=['$I5="Rural/Acreage"'],
                    fill=yellow_fill)
    )

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save("SCG_CMA_System_v7.xlsx")
print("Step 5 complete — Tab 1 cols 40-51 (AN-AY) built.")
print("  RURAL/ACREAGE banner AN2:AR2 (GOLD)")
print("  COMP SEARCH PARAMETERS banner AS2:AY2 (NAVY)")
print("  Acreage Band formula seeded in AN5, GREEN_BG.")
print("  3 dropdowns: Land Type, Primary Value Driver, Water Features.")
print("  1 dropdown: Requested Starting Radius (AW).")
print("  LIGHT_BLUE always-on CF for State col (AS).")
print("  CF on AN:AR — grey for In-Town, yellow for Rural.")
